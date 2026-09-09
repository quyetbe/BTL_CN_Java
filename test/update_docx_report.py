# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import shutil

print('Starting BTL_CN_Java.docx comprehensive update...')

# Backup original
shutil.copyfile('BTL_CN_Java.docx', 'BTL_CN_Java_Backup_BeforeAll.docx')

doc = docx.Document('BTL_CN_Java.docx')

# --- 1. Replace existing image blobs with new high-res screenshots ---
image_map = {
    'rId14': 'screenshots/hinh_login_quick_roles.png',
    'rId15': 'screenshots/hinh_trangchu_moi.png',
    'rId16': 'screenshots/hinh_tkb_grid_centered_scroll.png',
    'rId17': 'screenshots/hinh_tkb_pagination.png',
    'rId18': 'screenshots/hinh_3_5_xeplich_dialog.png',
    'rId19': 'screenshots/hinh_3_6_phonghoc.png',
    'rId20': 'screenshots/hinh_3_7_giangvien.png',
    'rId21': 'screenshots/hinh_danhmuc_monhoc.png',
    'rId22': 'screenshots/hinh_lophoc_vietnamese_centered.png',
    'rId23': 'screenshots/hinh_thongke_giangvien.png',
    'rId24': 'screenshots/hinh_taikhoan_6roles.png',
}

for rId, path in image_map.items():
    if rId in doc.part.related_parts and os.path.exists(path):
        part = doc.part.related_parts[rId]
        with open(path, 'rb') as f:
            part._blob = f.read()
        print(f'Replaced image blob for {rId} ({part.partname}) from {path}')

# --- 2. Update Section 2.1.1 (Actors & RBAC 6 roles) ---
for i, p in enumerate(doc.paragraphs):
    if '2.1.1. Tác nhân hệ thống (Actors)' in p.text:
        doc.paragraphs[i + 1].text = (
            'Hệ thống Quản lý Thời khóa biểu và Tài nguyên Phòng học (CNJ56) được thiết kế phục vụ 6 nhóm tác nhân chính với phân quyền chặt chẽ:\n'
            '1. Quản trị viên (ADMIN): Quản lý toàn diện tài khoản người dùng, phân cấp quyền hạn, đặt lại mật khẩu, theo dõi giám sát toàn bộ Nhật ký kiểm toán hệ thống (Audit Log).\n'
            '2. Ban Giám Hiệu (BAN_GIAM_HIEU): Xem Dashboard tổng quan, tra cứu TKB toàn trường, thống kê hiệu suất sử dụng tài nguyên, phê duyệt cấp 2 đối với các đơn xin đổi lịch dạy của giảng viên.\n'
            '3. Cán bộ Phòng Đào Tạo (PHONG_DAO_TAO): Lập lịch thời khóa biểu, quản lý danh mục phòng học, điều phối tài nguyên, xử lý xung đột lịch học, phê duyệt đổi lịch và xuất báo cáo.\n'
            '4. Trưởng Khoa / Trưởng Bộ Môn (TRUONG_KHOA): Quản lý hồ sơ giảng viên, danh mục môn học thuộc khoa, thẩm định và phê duyệt sơ bộ (cấp 1) các đơn đề xuất đổi lịch dạy.\n'
            '5. Giảng viên (GIANG_VIEN): Tra cứu thời khóa biểu giảng dạy cá nhân, tạo phiếu đề xuất xin đổi lịch dạy khi có việc đột xuất và theo dõi tiến độ phê duyệt của Khoa và Đào tạo.\n'
            '6. Sinh viên (SINH_VIEN): Tra cứu thời khóa biểu lớp học theo tuần, xem khung chương trình đào tạo chuẩn 130 tín chỉ 4 năm và tiến độ học tập cá nhân.'
        )
        print('Updated Section 2.1.1 with 6 RBAC actors!')
        break

# --- 3. Insert Hinh 2.1 image into Paragraph 96 if empty ---
p96 = doc.paragraphs[96]
if not p96.text.strip() and os.path.exists('docs/images/hinh_2_1_bieudo_phancap_chucnang.png'):
    p96.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run96 = p96.add_run()
    run96.add_picture('docs/images/hinh_2_1_bieudo_phancap_chucnang.png', width=Inches(6.2))
    print('Inserted Hinh 2.1 Draw.io picture into Paragraph 96!')

# --- 4. Update Captions in Section 3.4 ---
captions_update = {
    294: 'Hình 3.1: Giao diện Đăng nhập hệ thống & Quick Login 6 vai trò (LoginForm)',
    299: 'Hình 3.2: Giao diện Trang chủ tổng quan Dashboard 6 thẻ KPI (DashboardPanel)',
    303: 'Hình 3.3: Giao diện Lưới Thời Khóa Biểu Ma trận đa chiều (TimetableGridPanel - Matrix View)',
    307: 'Hình 3.4: Giao diện Quản lý Danh sách Thời Khóa Biểu & Phân trang (ThoiKhoaBieuPanel)',
    325: 'Hình 3.8: Giao diện Quản lý Môn học chuẩn khung CTĐT (MonHocPanel)',
    329: 'Hình 3.9: Giao diện Quản lý Lớp học sinh viên (LopHocPanel)',
    333: 'Hình 3.10: Giao diện Thống kê Giờ dạy & Tải đào tạo Giảng viên (ThongKePanel - Giảng viên)',
    338: 'Hình 3.11: Giao diện Quản trị Tài khoản người dùng & Phân quyền 6 vai trò RBAC (TaiKhoanPanel)',
}

for idx, text in captions_update.items():
    if idx < len(doc.paragraphs):
        doc.paragraphs[idx].text = text

# Update descriptions in Section 3.4
doc.paragraphs[295].text = 'Mô tả: Tiếp nhận tên đăng nhập và mật khẩu, kiểm tra trạng thái hoạt động của tài khoản trong CSDL qua AuthService (SHA-256). Tích hợp các nút chọn vai trò nhanh (Quick Login) hỗ trợ hội đồng nghiệm thu kiểm tra nhanh 6 vai trò phân quyền.'
doc.paragraphs[300].text = 'Mô tả: Hiển thị 6 thẻ thống kê KPI tổng quan (Phòng học, Giảng viên, Sinh viên, Môn học, Lịch xếp tuần, Yêu cầu chờ duyệt), biểu đồ phân bổ trạng thái và khối thông tin hệ thống chuyên nghiệp đặt gọn gàng ở phía dưới.'
doc.paragraphs[304].text = 'Mô tả: Trực quan hóa thời khóa biểu theo ma trận 12 tiết học x 7 ngày trong tuần. Hỗ trợ lọc xem theo Lớp học, theo Phòng học, theo Giảng viên hoặc Toàn trường; nhấp đúp vào ô để xem chi tiết ca học.'
doc.paragraphs[308].text = 'Mô tả: Danh sách bảng dữ liệu đầy đủ các lịch học với 2 hàng bộ lọc đa tiêu chí (Kỳ, Năm, Tuần, Thứ, Phòng, GV, Lớp), nút Xuất Excel/CSV và thanh phân trang PaginationBar điều hướng linh hoạt.'
doc.paragraphs[334].text = 'Mô tả: Báo cáo đo lường chi tiết tổng số giờ dạy, số tiết học và tải giảng dạy của từng giảng viên theo tuần và học kỳ, hỗ trợ phòng Đào tạo cân đối khối lượng công việc.'
doc.paragraphs[339].text = 'Mô tả: Phân hệ dành riêng cho ADMIN: quản lý tài khoản người dùng, cấp phát và thay đổi 6 vai trò RBAC (Admin, Ban Giám Hiệu, Phòng Đào Tạo, Trưởng Khoa, Giảng Viên, Sinh Viên), đặt lại mật khẩu và bật/tắt trạng thái khóa tài khoản.'

# --- 5. Insert New Sections 12 to 18 before 3.5 Kiem thu ---
p35 = None
for p in doc.paragraphs:
    if '3.5' in p.text and 'Kiểm thử' in p.text:
        p35 = p
        break

if p35:
    print('Inserting new sections before 3.5...')
    new_sections = [
        ('12. Màn hình Lưới Thời Khóa Biểu Dạng Khối (Block View) & TKB Giảng Viên',
         'screenshots/hinh_grid_block_schedule.png',
         'Hình 3.12: Giao diện Lưới Thời Khóa Biểu dạng Khối trực quan (Block View)',
         'Mô tả: Chế độ hiển thị dạng Khối (Block View) trực quan hóa các ca học theo từng khối thời gian sáng và chiều, giúp giảng viên và người quản lý dễ dàng nhận biết các khoảng trống phòng học và thời gian biểu trong tuần.',
         'screenshots/hinh_tkb_giangvien.png',
         'Hình 3.13: Giao diện Thời Khóa Biểu Giảng Dạy Cá Nhân của Giảng Viên',
         'Mô tả: Phân hệ cho phép giảng viên tra cứu lịch giảng dạy cá nhân trong từng học kỳ và tuần học cụ thể, hiển thị đầy đủ thông tin môn học, lớp sinh viên, phòng học và thời gian bắt đầu - kết thúc.'),

        ('13. Phân hệ Đề Xuất & Phê Duyệt Đổi Lịch 2 Cấp Độ',
         'screenshots/hinh_dexuat_doilich.png',
         'Hình 3.14: Giao diện Giảng viên Đề xuất Đổi lịch dạy và Theo dõi trạng thái đơn',
         'Mô tả: Giảng viên gửi đơn đề xuất đổi lịch dạy (chọn phòng mới, thứ mới, tiết mới). Hệ thống gọi XepLichService kiểm tra xung đột tức thời trước khi lưu đơn. Đơn chuyển qua quy trình duyệt 2 cấp: Trưởng Khoa duyệt cấp 1 -> Phòng Đào tạo / Ban Giám Hiệu duyệt chốt cấp 2 và tự động cập nhật TKB.'),

        ('14. Phân hệ Quản Lý Khung Chương Trình Đào Tạo 130 Tín Chỉ (CurriculumPanel)',
         'screenshots/hinh_ctdt_k24_130tc.png',
         'Hình 3.15: Giao diện Quản lý Khung Chương trình đào tạo 130 tín chỉ 4 năm (CurriculumPanel)',
         'Mô tả: Quản lý chi tiết toàn bộ 42 môn học với tổng cộng 130 tín chỉ phân bổ chuẩn hóa qua 8 học kỳ (4 năm học). Hỗ trợ lọc theo từng khóa sinh viên (K21, K22, K23, K24, K25) và tính toán tổng số tín chỉ tích lũy.'),

        ('15. Phân hệ Quản Lý Sinh Viên & Thanh Phân Trang Chuẩn Hóa (SinhVienPanel & PaginationBar)',
         'screenshots/hinh_sinhvien_pagination.png',
         'Hình 3.16: Giao diện Quản lý Sinh viên tích hợp Thanh Phân trang chuẩn hóa (PaginationBar)',
         'Mô tả: Quản lý danh mục 2.000 sinh viên chia theo các lớp và niên khóa. Tích hợp thanh phân trang PaginationBar (20 sinh viên/trang) với các nút Đầu, Trước, danh sách số trang, Sau, Cuối cùng chỉ báo Trang hiện tại / Tổng số trang và Tổng số bản ghi.'),

        ('16. Phân hệ Nhật Ký Kiểm Toán An Ninh Hệ Thống (AuditLogPanel)',
         'screenshots/hinh_audit_log.png',
         'Hình 3.17: Giao diện Nhật ký Kiểm toán Hoạt động & Bảo mật Hệ thống (AuditLogPanel)',
         'Mô tả: Giám sát toàn bộ hoạt động trong hệ thống. Ghi nhận chi tiết: Mã người dùng, Tên đăng nhập, Hành động (Đăng nhập, Xếp lịch, Đổi phòng, Sửa điểm...), Đối tượng tác động, Dữ liệu cũ -> Dữ liệu mới, Địa chỉ IP và Dấu thời gian chính xác tới từng giây.'),

        ('17. Phân hệ Thống Kê Phân Bổ Sinh Viên Theo Niên Khóa & Lớp Học',
         'screenshots/hinh_thongke_sinhvien.png',
         'Hình 3.18: Giao diện Thống kê Phân bổ Sinh viên theo Niên khóa và Lớp học',
         'Mô tả: Thống kê số lượng sinh viên theo từng khóa học (K21 - K25) và từng lớp chuyên ngành, tính toán tỷ lệ phân bổ phần trăm giúp nhà trường có cái nhìn tổng quan về quy mô đào tạo từng niên khóa.'),

        ('18. Phân hệ Phân Quyền Đa Cấp RBAC 6 Vai Trò & Chuyển Đổi Nhanh',
         'screenshots/hinh_login_quick_roles.png',
         'Hình 3.19: Cơ chế Đăng nhập nhanh và phân quyền 6 vai trò người dùng (RBAC)',
         'Mô tả: Hệ thống thiết lập ma trận phân quyền 6 vai trò người dùng chuyên biệt: ADMIN, BAN_GIAM_HIEU, PHONG_DAO_TAO, TRUONG_KHOA, GIANG_VIEN, SINH_VIEN. Menu điều hướng bên trái (Sidebar) tự động ẩn/hiện đúng phạm vi quyền hạn của từng vai trò.')
    ]

    for item in new_sections:
        title = item[0]
        img1, cap1, desc1 = item[1], item[2], item[3]
        
        p_title = p35.insert_paragraph_before(title)
        p_title.runs[0].bold = True
        p_title.runs[0].font.size = Pt(13)
        
        if os.path.exists(img1):
            p_img1 = p35.insert_paragraph_before()
            p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img1.add_run().add_picture(img1, width=Inches(5.8))
            
        p_cap1 = p35.insert_paragraph_before(cap1)
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.runs[0].italic = True
        p_cap1.runs[0].font.size = Pt(10.5)
        
        p_desc1 = p35.insert_paragraph_before(desc1)
        p_desc1.runs[0].font.size = Pt(11)
        
        if len(item) > 4:
            img2, cap2, desc2 = item[4], item[5], item[6]
            if os.path.exists(img2):
                p_img2 = p35.insert_paragraph_before()
                p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img2.add_run().add_picture(img2, width=Inches(5.8))
            p_cap2 = p35.insert_paragraph_before(cap2)
            p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap2.runs[0].italic = True
            p_cap2.runs[0].font.size = Pt(10.5)
            p_desc2 = p35.insert_paragraph_before(desc2)
            p_desc2.runs[0].font.size = Pt(11)

# --- 6. Update DANH MUC HINH ANH ---
list_figures = [
    'Hình 2.1 Biểu đồ phân cấp chức năng hệ thống (BFD) [Thiết kế trên Draw.io]	7',
    'Hình 2.2 Sơ đồ kiến trúc hệ thống phân tầng tích hợp Nhật ký hệ thống (Audit Log)	9',
    'Hình 2.3 Mô hình liên kết thực thể (ERD)	16',
    'Hình 2.4 Mô hình vật lý cơ sở dữ liệu (Physical Data Model)	17',
    'Hình 3.1: Giao diện Đăng nhập hệ thống & Quick Login 6 vai trò (LoginForm)	31',
    'Hình 3.2: Giao diện Trang chủ tổng quan Dashboard 6 thẻ KPI (DashboardPanel)	32',
    'Hình 3.3: Giao diện Lưới Thời Khóa Biểu Ma trận đa chiều (TimetableGridPanel - Matrix View)	32',
    'Hình 3.4: Giao diện Quản lý Danh sách Thời Khóa Biểu & Phân trang (ThoiKhoaBieuPanel)	33',
    'Hình 3.5: Hộp thoại Xếp lịch mới & Cảnh báo phát hiện Xung đột (XepLichDialog)	34',
    'Hình 3.6: Giao diện Quản lý Phòng học & Tài nguyên (PhongHocPanel)	34',
    'Hình 3.7: Giao diện Quản lý Giảng viên (GiangVienPanel)	35',
    'Hình 3.8: Giao diện Quản lý Môn học chuẩn khung CTĐT (MonHocPanel)	36',
    'Hình 3.9: Giao diện Quản lý Lớp học sinh viên (LopHocPanel)	36',
    'Hình 3.10: Giao diện Thống kê Giờ dạy & Tải đào tạo Giảng viên (ThongKePanel)	37',
    'Hình 3.11: Giao diện Quản trị Tài khoản người dùng & Phân quyền 6 vai trò (TaiKhoanPanel)	38',
    'Hình 3.12: Giao diện Lưới Thời Khóa Biểu dạng Khối trực quan (Block View)	39',
    'Hình 3.13: Giao diện Thời Khóa Biểu Giảng Dạy Cá Nhân của Giảng Viên	39',
    'Hình 3.14: Giao diện Giảng viên Đề xuất Đổi lịch dạy và Theo dõi trạng thái đơn	40',
    'Hình 3.15: Giao diện Quản lý Khung Chương trình đào tạo 130 tín chỉ 4 năm (CurriculumPanel)	41',
    'Hình 3.16: Giao diện Quản lý Sinh viên tích hợp Thanh Phân trang chuẩn hóa (PaginationBar)	42',
    'Hình 3.17: Giao diện Nhật ký Kiểm toán Hoạt động & Bảo mật Hệ thống (AuditLogPanel)	43',
    'Hình 3.18: Giao diện Thống kê Phân bổ Sinh viên theo Niên khóa và Lớp học	44',
    'Hình 3.19: Cơ chế Đăng nhập nhanh và phân quyền 6 vai trò người dùng (RBAC)	45',
]

start_fig = -1
end_fig = -1
for i, p in enumerate(doc.paragraphs[:60]):
    if 'DANH MỤC HÌNH ẢNH' in p.text:
        start_fig = i + 1
    if 'DANH MỤC BẢNG BIỂU' in p.text:
        end_fig = i
        break

if start_fig != -1 and end_fig != -1:
    print(f'Updating DANH MUC HINH ANH from P{start_fig} to P{end_fig}...')
    existing_count = end_fig - start_fig
    for j, fig_text in enumerate(list_figures):
        if j < existing_count:
            doc.paragraphs[start_fig + j].text = fig_text
        else:
            doc.paragraphs[end_fig].insert_paragraph_before(fig_text)

# --- 7. Update DANH MUC BANG BIEU ---
list_tables = [
    'Bảng 1.1: Bảng phân công nhiệm vụ thực hiện đề tài (Nhóm 2 thành viên)	5',
    'Bảng 1.2: Ma trận tiến độ phối hợp triển khai 10 tuần	6',
    'Bảng 2.1: Bảng ma trận phân quyền 6 vai trò người dùng (RBAC Matrix)	9',
    'Bảng 2.2: Mô tả màn hình Đăng nhập (LoginForm)	10',
    'Bảng 2.3: Mô tả màn hình Trang chủ Dashboard (DashboardPanel)	10',
    'Bảng 2.4: Mô tả màn hình Lưới Thời Khóa Biểu (TimetableGridPanel)	11',
    'Bảng 2.5: Mô tả màn hình Xếp lịch TKB (ThoiKhoaBieuPanel & XepLichDialog)	12',
    'Bảng 2.6: Mô tả màn hình Đề xuất & Phê duyệt đổi lịch 2 cấp (DeXuatDoiLichPanel & DoiLichPanel)	13',
    'Bảng 2.7: Mô tả màn hình Khung Chương trình đào tạo 130 tín chỉ (CurriculumPanel)	13',
    'Bảng 2.8: Mô tả màn hình Quản lý Sinh viên & Phân trang (SinhVienPanel & PaginationBar)	14',
    'Bảng 2.9: Mô tả màn hình Nhật ký Kiểm toán hệ thống (AuditLogPanel)	14',
    'Bảng 2.10: Mô tả màn hình Thống kê & Tra cứu phòng trống (ThongKePanel)	15',
    'Bảng 2.11: Bảng mô tả chi tiết vai trò từng file mã nguồn trong dự án	25',
    'Bảng 3.1: Bảng kịch bản và kết quả kiểm thử hệ thống (Test Cases Result)	48',
]

start_tbl = -1
end_tbl = -1
for i, p in enumerate(doc.paragraphs[:60]):
    if 'DANH MỤC BẢNG BIỂU' in p.text:
        start_tbl = i + 1
    if 'LỜI MỞ ĐẦU' in p.text:
        end_tbl = i
        break

if start_tbl != -1 and end_tbl != -1:
    print(f'Updating DANH MUC BANG BIEU from P{start_tbl} to P{end_tbl}...')
    existing_count = end_tbl - start_tbl
    for j, tbl_text in enumerate(list_tables):
        if j < existing_count:
            doc.paragraphs[start_tbl + j].text = tbl_text
        else:
            doc.paragraphs[end_tbl].insert_paragraph_before(tbl_text)

# Save updated docx
doc.save('BTL_CN_Java.docx')
print('Successfully saved BTL_CN_Java.docx!')

# Sync to other copies
for copy_name in ['BTL_CN_Java_Final.docx', 'BTL_CN_Java_Updated.docx', 'BTL_CN_Java_BW.docx']:
    shutil.copyfile('BTL_CN_Java.docx', copy_name)
    print(f'Synced to {copy_name}')

print('ALL DOCX UPDATES COMPLETED!')
