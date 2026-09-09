import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import zipfile
import shutil
import os

print("=== BẮT ĐẦU CẬP NHẬT BÁO CÁO (DOCX & MARKDOWN) ===")

src_docx = "BTL_CN_Java.docx"
target_docx = "BTL_CN_Java_Updated.docx"
new_img_path = "docs/images/hinh_2_2_kientruc_hethong_moi.png"

# 1. Tạo tệp docx mới với ảnh Draw.io thay thế
with zipfile.ZipFile(src_docx, 'r') as zin:
    with zipfile.ZipFile(target_docx, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == 'word/media/image17.png':
                with open(new_img_path, 'rb') as f_img:
                    zout.writestr(item, f_img.read())
                print("✅ Đã cập nhật ảnh word/media/image17.png thành sơ đồ Draw.io mới!")
            else:
                zout.writestr(item, zin.read(item.filename))

# 2. Đọc và chỉnh sửa nội dung văn bản trong target_docx
doc = docx.Document(target_docx)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# A. Cập nhật Bảng 2 (Table 2)
table2 = doc.tables[2]
for row in table2.rows:
    for cell in row.cells:
        if '[Họ và tên thành viên 2]' in cell.text:
            cell.text = cell.text.replace('[Họ và tên thành viên 2]', 'Hà Thái Bảo (Thành viên)')
            print("✅ Đã cập nhật tên Hà Thái Bảo vào Table 2!")

# B. Tìm vị trí các đoạn văn
idx_13 = -1
idx_14 = -1
idx_22 = -1
idx_hinh22 = -1

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text == '1.3 Các công nghệ sử dụng.':
        idx_13 = i
    elif text == '1.4 Kết chương.':
        idx_14 = i
    elif '2.2 Kiến trúc hệ thống phân tầng' in text:
        idx_22 = i
    elif 'Hình 2.2' in text and i > 50:
        idx_hinh22 = i

print(f"Vị trí tìm thấy: 1.3 -> {idx_13}, 1.4 -> {idx_14}, 2.2 -> {idx_22}, Hình 2.2 -> {idx_hinh22}")

# C. Cập nhật Mục 2.2 Kiến trúc hệ thống: Thêm ghi chú Draw.io
if idx_22 != -1:
    p_next = doc.paragraphs[idx_22 + 1]
    note_drawio = (
        " (Ghi chú: Toàn bộ sơ đồ kiến trúc phân tầng được thiết kế chi tiết bằng công cụ Draw.io (diagrams.net), "
        "tệp mã nguồn sơ đồ được lưu trữ tại docs/architecture/kientruc_hethong_phantang.drawio để tiện theo dõi và hiệu chỉnh)."
    )
    if "Draw.io" not in p_next.text:
        p_next.text = p_next.text + note_drawio
        print("✅ Đã bổ sung ghi chú Draw.io vào Mục 2.2!")

if idx_hinh22 != -1:
    p_caption = doc.paragraphs[idx_hinh22]
    p_caption.text = "Hình 2.2: Sơ đồ kiến trúc hệ thống phân tầng tích hợp Nhật ký hệ thống (Audit Log) [Thiết kế trên Draw.io]"
    print("✅ Đã cập nhật chú thích Hình 2.2 có [Thiết kế trên Draw.io]!")

# D. Cập nhật mục 1.3, 1.4, 1.5, 1.6
if idx_14 != -1:
    doc.paragraphs[idx_14].text = "1.6 Kết chương."

p_ketchuong = doc.paragraphs[idx_14]

p_13 = doc.paragraphs[idx_13]
p_13.text = "1.3 Các công nghệ và công cụ sử dụng."
p_13_desc = p_13.insert_paragraph_before(
    "Đề tài áp dụng các công nghệ tiêu chuẩn trong hệ sinh thái Java doanh nghiệp:\n"
    "- Ngôn ngữ lập trình: Java SE (phiên bản 17 LTS trở lên), biên dịch và chạy trên nền tảng NetBeans IDE / IntelliJ IDEA.\n"
    "- Giao diện người dùng: Thư viện đồ họa Java Swing kết hợp gói tùy biến giao diện phẳng Flat UI (khử răng cưa khử mờ font tiếng Việt Segoe UI, tối ưu hiển thị Full màn hình MAXIMIZED_BOTH).\n"
    "- Kết nối Cơ sở dữ liệu: Trình điều khiển MySQL Connector/J 8.3.0, thực thi 100% qua JDBC PreparedStatement chống SQL Injection.\n"
    "- Hệ quản trị CSDL: MySQL Server 8.0 chạy trên môi trường XAMPP Localhost cổng mặc định 3306.\n"
    "- Thiết kế sơ đồ kiến trúc & CSDL: Sử dụng phần mềm Draw.io (diagrams.net) để mô hình hóa kiến trúc phân tầng chuẩn mực.\n"
    "- Cơ chế an toàn & Mã hóa: Thuật toán băm mật khẩu SHA-256 an toàn cao, Interceptor bắt sự kiện tự động phục vụ Audit Log."
)

p_14_heading = p_ketchuong.insert_paragraph_before("1.4 Kế hoạch thực hiện đề tài.")
p_14_desc = p_ketchuong.insert_paragraph_before(
    "Đề tài được tiến hành trong thời gian 10 tuần với các giai đoạn nối tiếp khoa học: "
    "Tuần 1-2: Khảo sát & Phân tích yêu cầu nghiệp vụ; "
    "Tuần 3-4: Thiết kế CSDL 3NF & Thiết kế sơ đồ kiến trúc Draw.io; "
    "Tuần 5-6: Lập trình Backend, CSDL & Thuật toán kiểm tra xung đột thời khóa biểu 8 ràng buộc; "
    "Tuần 7-8: Lập trình Giao diện Java Swing Flat UI, Lưới TKB tuần & Phân trang danh mục; "
    "Tuần 9: Tích hợp hệ thống, Kiểm thử tự động 15 Test Case & Tối ưu hiệu năng; "
    "Tuần 10: Đóng gói sản phẩm, hoàn thiện báo cáo và bảo vệ đồ án."
)

p_15_heading = p_ketchuong.insert_paragraph_before("1.5 Phân công nhiệm vụ thực hiện đề tài (Nhóm 2 thành viên).")

p_15_desc1 = p_ketchuong.insert_paragraph_before(
    "Để bảo đảm tiến độ và chất lượng sản phẩm phần mềm, Nhóm 11 (gồm 2 thành viên) đã tiến hành phân chia nhiệm vụ "
    "rõ ràng, tương xứng với năng lực chuyên môn của từng sinh viên theo mô hình phân tầng phát triển (Backend - Architecture & Frontend - QA):"
)

# Chèn bảng phân công chi tiết 1.5
table_15 = doc.add_table(rows=3, cols=6)
table_15.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ["STT", "Họ và tên", "Mã sinh viên", "Lớp hành chính", "Vai trò đảm nhiệm", "Nhiệm vụ chính & Tỷ lệ đóng góp"]
hdr_cells = table_15.rows[0].cells
for j, h in enumerate(headers):
    hdr_cells[j].text = h
    set_cell_background(hdr_cells[j], "1E293B")
    for p in hdr_cells[j].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)

# Dữ liệu SV 1: Nguyễn Mạnh Quyết
row1 = table_15.rows[1].cells
row1[0].text = "1"
row1[1].text = "Nguyễn Mạnh Quyết\n(Trưởng nhóm)"
row1[2].text = "20231053"
row1[3].text = "DCCNTT.14.3"
row1[4].text = "Trưởng nhóm &\nBackend Developer /\nKiến trúc sư hệ thống"
row1[5].text = (
    "• Lập kế hoạch dự án, phân rã chức năng nghiệp vụ đề tài CNJ56.\n"
    "• Thiết kế mô hình dữ liệu quan hệ chuẩn 3NF và viết kịch bản database.sql (11 bảng, ràng buộc toàn vẹn).\n"
    "• Thiết kế sơ đồ kiến trúc hệ thống phân tầng trên công cụ Draw.io (kientruc_hethong_phantang.drawio).\n"
    "• Xây dựng tầng kết nối DBConnection (Singleton Pool) và toàn bộ 10 lớp DAO (JDBC PreparedStatement chống SQLi).\n"
    "• Nghiên cứu và cài đặt thuật toán kiểm tra xung đột thời khóa biểu đa chiều 8 ràng buộc (XepLichService).\n"
    "• Xây dựng cơ chế xác thực băm mật khẩu SHA-256 và phân quyền RBAC 6 vai trò (AuthService).\n"
    "• Viết bộ kiểm thử tự động 15 ca kiểm thử logic và hiệu năng (TestLogicRunner.java).\n"
    "➔ Mức độ hoàn thành: 100% | Tỷ lệ đóng góp: 50%"
)

# Dữ liệu SV 2: Hà Thái Bảo
row2 = table_15.rows[2].cells
row2[0].text = "2"
row2[1].text = "Hà Thái Bảo\n(Thành viên)"
row2[2].text = "20231045"
row2[3].text = "DCCNTT.14.3"
row2[4].text = "Thành viên &\nFrontend Developer /\nKiểm thử phần mềm (QA)"
row2[5].text = (
    "• Thiết kế và xây dựng toàn bộ giao diện Desktop Java Swing theo phong cách Flat UI hiện đại.\n"
    "• Xây dựng MainForm với thanh điều hướng thích ứng phân quyền, tự động phóng to toàn màn hình.\n"
    "• Xây dựng ma trận Lưới thời khóa biểu tuần trực quan (TimetableGridPanel) tích hợp tính năng thu phóng Zoom.\n"
    "• Xây dựng toàn bộ các Form danh mục và hộp thoại (ThoiKhoaBieuPanel, XepLichDialog, MonHocPanel, GiangVienPanel, SinhVienPanel, PhongHocPanel, LopHocPanel, CurriculumPanel, DeXuatDoiLichPanel, AuditLogPanel, TaiKhoanPanel).\n"
    "• Xây dựng component dùng chung thanh phân trang PaginationBar (20 mục/trang).\n"
    "• Cài đặt ThongKeService (thống kê tải GV, quét phòng trống) và ExportUtil (xuất CSV/Excel UTF-8 BOM).\n"
    "• Thiết kế kịch bản kiểm thử giao diện người dùng, kiểm tra biên và hoàn thiện tài liệu báo cáo đồ án.\n"
    "➔ Mức độ hoàn thành: 100% | Tỷ lệ đóng góp: 50%"
)

for row_idx in [1, 2]:
    r_cells = table_15.rows[row_idx].cells
    bg = "F8FAFC" if row_idx == 1 else "FFFFFF"
    for c_idx, cell in enumerate(r_cells):
        set_cell_background(cell, bg)
        set_cell_margins(cell, 80, 80, 100, 100)
        for p in cell.paragraphs:
            if c_idx in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9.5)

p_ketchuong._p.addprevious(table_15._tbl)

p_detail = p_ketchuong.insert_paragraph_before(
    "\nĐánh giá chung về kết quả phối hợp của nhóm 2 thành viên:\n"
    "- Cả hai sinh viên Nguyễn Mạnh Quyết và Hà Thái Bảo đều chủ động, bám sát kế hoạch tiến độ 10 tuần, phối hợp nhịp nhàng giữa tầng Backend (Xử lý thuật toán, CSDL) và tầng Frontend (Giao diện Java Swing, Trực quan hóa).\n"
    "- Toàn bộ 15/15 bài kiểm thử Unit Test đều vượt qua thành công, hệ thống vận hành trơn tru với dữ liệu thực nghiệm 690 lịch học, 2.000 sinh viên, 506 giảng viên, 40 lớp học và 20 phòng học, không xảy ra bất kỳ xung đột nào.\n"
    "- Mức độ hoàn thành của cả hai thành viên đạt 100% khối lượng công việc được giao với tỷ lệ đóng góp đồng đều 50% - 50%."
)

doc.save(target_docx)
print(f"✅ Đã lưu tệp docx cập nhật thành công: {target_docx}")

# Thử ghi đè BTL_CN_Java.docx nếu Word không khóa
try:
    shutil.copy2(target_docx, src_docx)
    print(f"✅ Đã ghi đè trực tiếp vào tệp gốc: {src_docx}")
except Exception as e:
    print(f"ℹ️ Tệp {src_docx} đang được mở trong Word bởi bạn. Bản mới đã được lưu tại {target_docx} (bạn có thể đóng Word và ghi đè bất cứ lúc nào).")
