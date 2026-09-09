import docx
import re

doc = docx.Document('BTL_CN_Java.docx')
for i, p in enumerate(doc.paragraphs):
    if 'Hình 2.2' in p.text:
        print(f"Paragraph {i}: {p.text}")
        for j in range(max(0, i - 2), min(len(doc.paragraphs), i + 6)):
            xml = doc.paragraphs[j]._element.xml
            embeds = re.findall(r'r:embed="([^"]+)"', xml)
            has_draw = "w:drawing" in xml
            print(f"  p[{j}] (len={len(doc.paragraphs[j].text)}): text='{doc.paragraphs[j].text[:40]}' draws={has_draw} embeds={embeds}")
